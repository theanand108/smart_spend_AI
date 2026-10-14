from io import BytesIO

from openpyxl import Workbook

from src.statement_importer import parse_google_pay_text, parse_phonepe_csv, parse_paytm_xlsx, parse_statement


PHONEPE_SAMPLE = '''Transaction Statement for 6307183070
Duration,"30 Jul, 2026 - 29 Aug, 2026"

Date,Time,Transaction Details,Transaction ID,UTR,Transaction Type,Credit/debit instrument,Amount
"Aug 29, 2026","12:30 am","Paid to Vishal Bhai 2","T2608290030197862671909","653919911522","DEBIT","Paid by XXXXXXX1831","100"
"Aug 26, 2026","09:00 am","Received from MUKUND KUSHWAHA","T2608260900463826377542","660440391566","CREDIT","Credited to XXXXXXX1831","5000"
"Aug 25, 2026","03:31 pm","Paid to Anand bank","T2608251531164719873160","515202082686","DEBIT","Paid by XXXXXXX1831","299"
'''


GPAY_SAMPLE = '''Transaction statement
Date & time Transaction details Amount
01 Jul, 2026
02:20 PM
Paid to mukund Kushwaha
UPI Transaction ID: 618266600033
Paid by Uttar Pradesh Gramin Bank 2909
₹1,000
01 Jul, 2026
04:09 PM
Paid to EKART
UPI Transaction ID: 654883594580
Paid by Uttar Pradesh Gramin Bank 2909
₹369
02 Jul, 2026
05:00 PM
Received from MUKUND KUSHWAHA
UPI Transaction ID: 999999999999
Received by Uttar Pradesh Gramin Bank 2909
₹2,500
'''


def test_phonepe_csv_normalizes_debits_and_credits():
    result = parse_phonepe_csv(PHONEPE_SAMPLE)

    assert len(result.transactions) == 3
    assert len(result.expenses) == 2
    assert len(result.received) == 1

    assert result.expenses[0].merchant_name == "Vishal Bhai 2"
    assert result.expenses[0].amount == 100
    assert result.expenses[0].source_transaction_id == "T2608290030197862671909"
    assert result.expenses[0].note == "Paid to Vishal Bhai 2"

    assert result.received[0].merchant_name == "MUKUND KUSHWAHA"
    assert result.received[0].amount == 5000
    assert result.received[0].source_transaction_id == "T2608260900463826377542"
    assert result.received[0].note == "Received from MUKUND KUSHWAHA"


def test_phonepe_csv_accepts_september_abbreviation():
    csv_text = '''Transaction Statement for 6307183070
Duration,"26 Aug, 2026 - 25 Sept, 2026"

Date,Time,Transaction Details,Transaction ID,UTR,Transaction Type,Credit/debit instrument,Amount
"Sept 25, 2026","05:15 pm","Paid to NEERAJ KUMAR","T2609251715303336256820","073306645051","DEBIT","Paid by XXXXXXX1831","35"
"Sept 24, 2026","02:18 pm","Paid to Himanshu","R10502609241418466445214036","324041203358","DEBIT","Paid by XXXXXXX1831","175"
"Aug 26, 2026","09:00 am","Received from MUKUND KUSHWAHA","T2608260900463826377542","660440391566","CREDIT","Credited to XXXXXXX1831","5000"
'''
    result = parse_phonepe_csv(csv_text)

    assert len(result.transactions) == 3
    assert len(result.expenses) == 2
    assert len(result.received) == 1
    assert result.expenses[0].date.strftime("%Y-%m-%d %H:%M") == "2026-09-25 17:15"
    assert result.expenses[1].date.strftime("%Y-%m-%d %H:%M") == "2026-09-24 14:18"


def test_google_pay_pdf_text_normalizes_paid_and_received_transactions():
    result = parse_google_pay_text(GPAY_SAMPLE)

    assert len(result.transactions) == 3
    assert len(result.expenses) == 2
    assert len(result.received) == 1

    assert result.expenses[0].merchant_name == "mukund Kushwaha"
    assert result.expenses[0].amount == 1000
    assert result.expenses[0].source_transaction_id == "618266600033"
    assert result.expenses[0].note == "Paid to mukund Kushwaha"

    assert result.expenses[1].merchant_name == "EKART"
    assert result.expenses[1].note == "Paid to EKART"
    assert result.received[0].merchant_name == "MUKUND KUSHWAHA"
    assert result.received[0].amount == 2500
    assert result.received[0].note == "Received from MUKUND KUSHWAHA"


def test_gpay_converted_flat_csv_is_not_treated_as_a_supported_transaction_table():
    converted_csv = '''"Transactionstatement"\n"01Jul,2026"\n"02:20PM"\n"Paidto","EKART"\n"UPITransaction","ID:654883594580"\n"₹369"\n'''

    try:
        parse_phonepe_csv(converted_csv)
    except ValueError as exc:
        assert "header" in str(exc).lower()
    else:
        raise AssertionError("A flattened PDF-to-CSV file should not be accepted as a structured CSV.")



def _paytm_workbook_bytes():
    workbook = Workbook()
    workbook.active.title = "Summary"
    sheet = workbook.create_sheet("Passbook Payment History")
    sheet.append([
        "Date", "Time", "Transaction Details", "Other Transaction Details (UPI ID or A/c No)",
        "Your Account", "Amount", "UPI Ref No.", "Order ID", "Remarks", "Tags", "Comment",
    ])
    sheet.append(["17/04/2026", "19:09:56", "Received from One97 Communications Limited", "poweraccess.paytm3@axisbank", "Fino Payments Bank - 31", "+100.00", "827494321076", None, None, "# Money Received", None])
    sheet.append(["16/04/2026", "22:10:12", "Cashback Received from One97 Communications Limited", "poweraccess.paytm3@axisbank", "Fino Payments Bank - 31", "+50.00", "650171861066", None, "Payment on Flipkart using UPI on Apr 01 2026", "# Cashback", None])
    sheet.append(["04/04/2026", "18:24:41", "Cashback Received from One97 Communications Limited", "poweraccess.paytm3@axisbank", "Fino Payments Bank - 31", "+2.00", "189616340946", None, "Paying using UPI on Apr 04 2026", "# Cashback", None])
    sheet.append(["04/04/2026", "18:24:35", "Cashback Received from One97 Communications Limited", "poweraccess.paytm3@axisbank", "Fino Payments Bank - 31", "+1.00", "189593710946", None, "Paying using UPI on Apr 04 2026", "# Cashback", None])
    sheet.append(["04/04/2026", "18:23:48", "Recharge of Jio Mobile 6307183070", "paybil3066@ptybl on Paytm", "Fino Payments Bank - 31", "-299.00", "609442929881", "26848413096", None, "# Bill Payments", None])
    sheet.append(["04/04/2026", "02:04:27", "Paid to Blinkit", "blinkit.rzp@hdfcbank", "Fino Payments Bank - 31", "-383.00", "609442100191", None, "PayviaRazorpay", "# Groceries", None])
    sheet.append(["01/04/2026", "22:08:46", "Paid to Flipkart Payments", "paytm-56505013@ptybl on Paytm", "Fino Payments Bank - 31", "-186.00", "609169306754", None, None, "# Shopping", None])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_paytm_xlsx_imports_debits_credits_and_reference_ids():
    result = parse_paytm_xlsx(_paytm_workbook_bytes())

    assert len(result.transactions) == 7
    assert len(result.expenses) == 3
    assert len(result.received) == 4
    assert sum(t.amount for t in result.expenses) == 868
    assert sum(t.amount for t in result.received) == 153
    assert result.expenses[0].merchant_name == "Recharge of Jio Mobile 6307183070"
    assert result.expenses[0].source_transaction_id == "609442929881"
    assert result.expenses[0].source == "paytm"
    assert result.received[1].merchant_name == "Cashback Received from One97 Communications Limited"
    assert "Payment on Flipkart" in result.received[1].note


def test_parse_statement_routes_xlsx_to_paytm_parser():
    result = parse_statement("Paytm_UPI_Statement.xlsx", _paytm_workbook_bytes())
    assert len(result.transactions) == 7
    assert all(transaction.source == "paytm" for transaction in result.transactions)


def test_paytm_xlsx_rejects_unrelated_workbook():
    workbook = Workbook()
    workbook.active.append(["hello", "world"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    try:
        parse_paytm_xlsx(output.getvalue())
    except ValueError as exc:
        assert "Paytm" in str(exc)
    else:
        raise AssertionError("An unrelated workbook should not be accepted as a Paytm statement.")
